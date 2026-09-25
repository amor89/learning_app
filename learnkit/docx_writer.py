"""Render handbook blocks to a Word document with python-docx.

A4 pages with a wide right margin, so there is room to annotate.
Markdown blocks go through the same HTML renderer as the PDF, then a small
HTML-to-Word converter that handles the tags lesson content uses.
"""

from __future__ import annotations

from html.parser import HTMLParser
from typing import Any

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from learnkit.handbook import Block, Code, Contents, Heading, Items, Markdown, PageBreak
from learnkit.render import md_inline, md_to_html

MONO = "Consolas"
CODE_PT = 8
BODY_PT = 10.5
CODE_FILL = "F2F4F7"
MARK_FILL = "FFF4C2"


def shade(cell: Any, fill: str) -> None:
    """Set a table cell's background colour."""
    props = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    props.append(shd)


def mono_run(paragraph: Any, text: str, size: float = CODE_PT, bold: bool = False) -> Any:
    run = paragraph.add_run(text)
    run.font.name = MONO
    run.font.size = Pt(size)
    run.bold = bold
    run._element.rPr.rFonts.set(qn("w:eastAsia"), MONO)
    return run


def tight(paragraph: Any) -> None:
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.line_spacing = 1.0


class HtmlToDocx(HTMLParser):
    """Convert the HTML produced by ``md_to_html`` into paragraphs in ``container``."""

    def __init__(self, container: Any, list_style: str | None = None) -> None:
        super().__init__(convert_charrefs=True)
        self.container = container
        self.paragraph: Any = None
        self.bold = 0
        self.italic = 0
        self.code = 0
        self.lists: list[str] = []
        self.pre: list[str] | None = None
        self.table: list[list[str]] | None = None
        self.cell: list[str] | None = None
        self.default_style = list_style

    def _para(self, style: str | None = None) -> Any:
        self.paragraph = self.container.add_paragraph(style=style)
        return self.paragraph

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if self.pre is not None:
            return
        if tag == "p" and self.cell is None:
            self._para(self.default_style)
        elif tag in {"ul", "ol"}:
            self.lists.append("List Bullet" if tag == "ul" else "List Number")
        elif tag == "li":
            self._para(self.lists[-1] if self.lists else "List Bullet")
        elif tag in {"strong", "b"}:
            self.bold += 1
        elif tag in {"em", "i"}:
            self.italic += 1
        elif tag == "code":
            self.code += 1
        elif tag == "pre":
            self.pre = []
        elif tag == "table":
            self.table = []
        elif tag == "tr" and self.table is not None:
            self.table.append([])
        elif tag in {"td", "th"}:
            self.cell = []
        elif tag == "br" and self.paragraph is not None:
            self.paragraph.add_run().add_break()

    def handle_endtag(self, tag: str) -> None:
        if tag == "pre" and self.pre is not None:
            add_code(self.container, "".join(self.pre).rstrip("\n"), numbered=False)
            self.pre = None
            self.paragraph = None
        elif self.pre is not None:
            return
        elif tag in {"ul", "ol"} and self.lists:
            self.lists.pop()
        elif tag in {"strong", "b"}:
            self.bold -= 1
        elif tag in {"em", "i"}:
            self.italic -= 1
        elif tag == "code":
            self.code -= 1
        elif tag in {"td", "th"} and self.table is not None and self.cell is not None:
            self.table[-1].append("".join(self.cell).strip())
            self.cell = None
        elif tag == "table" and self.table is not None:
            add_table(self.container, self.table)
            self.table = None
        elif tag in {"p", "li"}:
            self.paragraph = None

    def handle_data(self, data: str) -> None:
        if self.pre is not None:
            self.pre.append(data)
            return
        if self.cell is not None:
            self.cell.append(data)
            return
        if not data.strip() and self.paragraph is None:
            return
        if self.paragraph is None:
            self._para(self.default_style)
        if self.code:
            mono_run(self.paragraph, data, size=BODY_PT - 1)
        else:
            run = self.paragraph.add_run(data)
            run.bold = self.bold > 0
            run.italic = self.italic > 0


def add_html(container: Any, markup: str, style: str | None = None) -> None:
    parser = HtmlToDocx(container, style)
    parser.feed(markup)
    parser.close()


def add_table(container: Any, rows: list[list[str]]) -> None:
    if not rows:
        return
    width = max(len(r) for r in rows)
    table = container.add_table(rows=len(rows), cols=width)
    table.style = "Table Grid"
    for r, row in enumerate(rows):
        for c, text in enumerate(row):
            cell = table.cell(r, c)
            cell.text = text
            if r == 0:
                for run in cell.paragraphs[0].runs:
                    run.bold = True


def add_code(
    container: Any, code: str, numbered: bool = True, notes: list[tuple[int, str]] | None = None
) -> None:
    """Code in a shaded one-cell table. Long lines wrap inside the cell."""
    notes = notes or []
    marks = {line: i for i, (line, _) in enumerate(notes, start=1)}
    table = container.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    cell = table.cell(0, 0)
    shade(cell, CODE_FILL)
    lines = code.rstrip("\n").split("\n")
    first = True
    for n, line in enumerate(lines, start=1):
        paragraph = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        tight(paragraph)
        if numbered:
            # Hanging indent: a wrapped line continues under the code, not the line number.
            paragraph.paragraph_format.left_indent = Cm(1.0)
            paragraph.paragraph_format.first_line_indent = Cm(-1.0)
            label = f"[{marks[n]}]" if n in marks else str(n)
            mono_run(paragraph, f"{label:>4}  ", bold=n in marks).font.color.rgb = RGBColor(
                0x5B, 0x68, 0x75
            )
        mono_run(paragraph, line if line else " ")
    if notes:
        for line_no, note in notes:
            paragraph = container.add_paragraph(style="List Number")
            paragraph.add_run(f"Line {line_no}. ").bold = True
            add_inline(paragraph, note)
    else:
        container.add_paragraph()


def add_inline(paragraph: Any, markdown_text: str) -> None:
    """Append inline Markdown (bold, italic, code) to an existing paragraph."""

    class Inline(HtmlToDocx):
        def _para(self, style: str | None = None) -> Any:
            self.paragraph = paragraph
            return paragraph

    parser = Inline(None)
    parser.paragraph = paragraph
    parser.feed(md_inline(markdown_text))
    parser.close()


def set_page(document: Any) -> None:
    section = document.sections[0]
    section.page_height = Cm(29.7)
    section.page_width = Cm(21.0)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(4.5)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(BODY_PT)


def render_docx(blocks: list[Block], title: str, path: str) -> None:
    """Write blocks to a .docx file."""
    document = Document()
    set_page(document)
    document.core_properties.title = title
    started = False
    for block in blocks:
        if isinstance(block, Heading):
            document.add_heading(block.text, level=min(block.level, 9) if block.level > 0 else 0)
        elif isinstance(block, Markdown):
            add_html(document, md_to_html(block.text))
        elif isinstance(block, Code):
            add_code(document, block.code, numbered=True, notes=block.notes)
        elif isinstance(block, Items):
            style = None if block.checklist else "List Number" if block.ordered else "List Bullet"
            for item in block.items:
                add_inline(document.add_paragraph(style=style), item)
        elif isinstance(block, PageBreak):
            if started:
                document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        elif isinstance(block, Contents):
            for level, text, _anchor in block.entries:
                paragraph = document.add_paragraph(text)
                paragraph.paragraph_format.left_indent = Cm(0.6 * (level - 1))
                tight(paragraph)
                if level == 1:
                    paragraph.runs[0].bold = True
        started = True
    document.save(path)

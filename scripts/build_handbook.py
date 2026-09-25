"""Build the handbook in PDF (A5, for iPhone) and Word (A4, for annotation).

Three levels: full handbook, one per track and one per module. Output goes to
site/downloads/ and is committed, so GitHub Pages serves it.

The build fails when a lesson in the app bundle is missing from the full handbook.

Run from the repository root: ``python scripts/build_handbook.py``.
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from docx import Document  # noqa: E402
from pygments.formatters import HtmlFormatter  # noqa: E402
from weasyprint import HTML  # noqa: E402

from learnkit.docx_writer import render_docx  # noqa: E402
from learnkit.handbook import Heading, build_blocks, render_html, scopes  # noqa: E402
from learnkit.loader import load_course  # noqa: E402
from learnkit.render import CSS_CLASS  # noqa: E402

DOWNLOADS = ROOT / "site" / "downloads"
LESSON_BUNDLE = ROOT / "site" / "content" / "lessons"

PDF_CSS = """
@page {
  size: A5;
  margin: 13mm 10mm 15mm 10mm;
  @bottom-center { content: counter(page); font: 7.5pt "DejaVu Sans"; color: #5b6875; }
}
html { font: 9pt/1.45 "DejaVu Sans", "Liberation Sans", sans-serif; color: #17212b; }
h1 { font-size: 16pt; color: #1f5f8b; margin: 0 0 8pt; }
h2 { font-size: 13pt; color: #1f5f8b; margin: 12pt 0 6pt; }
h3 { font-size: 11pt; margin: 14pt 0 4pt; border-bottom: 1px solid #d5dbe3; padding-bottom: 2pt; }
h4 { font-size: 9.5pt; margin: 10pt 0 3pt; color: #1f5f8b; }
h5 { font-size: 9pt; margin: 9pt 0 2pt; }
h3, h4, h5 { break-after: avoid; }
p { margin: 3pt 0; }
a { color: #1f5f8b; text-decoration: none; }
code { font: 7.8pt "DejaVu Sans Mono", monospace; background: #eef1f5; padding: 0 1.5pt; }
.pagebreak { break-before: page; }
ul.items, ol.items, ol.notes, ul, ol { margin: 3pt 0; padding-left: 13pt; }
li { margin: 1.5pt 0; }
ol.notes { font-size: 8.3pt; }
.note-line { font-weight: bold; }
table { border-collapse: collapse; width: 100%; margin: 4pt 0; }
th, td { border: 0.5pt solid #d5dbe3; padding: 2pt 4pt; font-size: 8pt; vertical-align: top; }
table.codetable {
  background: #f5f7fa; border: 0.5pt solid #d5dbe3; table-layout: fixed;
  font: 7.2pt/1.35 "DejaVu Sans Mono", monospace;
}
table.codetable td { border: 0; padding: 0 3pt; font-size: 7.2pt; }
table.codetable td.lnum { width: 16pt; text-align: right; color: #5b6875; }
/* Wrapped lines get a hanging indent, so a continuation never looks like a new line. */
table.codetable td.src {
  white-space: pre-wrap; overflow-wrap: break-word; padding-left: 14pt; text-indent: -11pt;
}
table.codetable tr.ann { background: #fff4c2; }
table.codetable tr { break-inside: avoid; }
.mark {
  display: inline-block; min-width: 9pt; border-radius: 5pt; text-align: center;
  background: #1f5f8b; color: #fff; font: bold 6.5pt "DejaVu Sans";
}
div.hl pre {
  background: #f5f7fa; border: 0.5pt solid #d5dbe3; padding: 4pt 5pt;
  font: 7.2pt/1.35 "DejaVu Sans Mono", monospace;
  white-space: pre-wrap; overflow-wrap: break-word;
}
ul.toc { list-style: none; padding: 0; }
ul.toc li { margin: 1pt 0; }
ul.toc a::after { content: leader('.') target-counter(attr(href), page); }
.toc1 { font-weight: bold; margin-top: 5pt !important; }
.toc2 { padding-left: 9pt; }
.toc3 { padding-left: 18pt; font-size: 8.3pt; }
div.hl pre code { background: none; padding: 0; font: inherit; }
ul.checklist { list-style: none; padding-left: 0; }
ul.checklist li { margin: 4pt 0; }
"""


def pdf_css() -> str:
    """Stylesheet plus Pygments colours."""
    styles = HtmlFormatter(style="friendly").get_style_defs(f".{CSS_CLASS}")
    return PDF_CSS + str(styles)


def docx_text(path: Path) -> str:
    """All paragraph and table text of a .docx, for the coverage check."""
    document = Document(str(path))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts += [cell.text for cell in row.cells]
    return "\n".join(parts)


def check_coverage(pdf_html: str, docx_path: Path) -> None:
    """Fail when any lesson in the app bundle is missing from the full handbook."""
    bundle_ids = sorted(p.stem for p in LESSON_BUNDLE.glob("*.json"))
    words = docx_text(docx_path)
    missing = [i for i in bundle_ids if f'id="lesson-{i}"' not in pdf_html or f"{i} " not in words]
    if missing:
        raise SystemExit(f"Lessons missing from the handbook: {missing}")
    print(
        f"Coverage: all {len(bundle_ids)} app lessons appear in the full handbook (PDF and Word)."
    )


def main() -> None:
    """Build every handbook scope."""
    course = load_course()
    DOWNLOADS.mkdir(parents=True, exist_ok=True)
    css = pdf_css()
    keep = set()
    for scope in scopes(course):
        blocks = build_blocks(scope, course)
        page = render_html(blocks, scope.title, css)
        pdf_path = DOWNLOADS / f"{scope.stem}.pdf"
        docx_path = DOWNLOADS / f"{scope.stem}.docx"
        HTML(string=page, base_url=str(ROOT)).write_pdf(str(pdf_path))
        render_docx(blocks, scope.title, str(docx_path))
        keep |= {pdf_path.name, docx_path.name}
        headings = sum(1 for b in blocks if isinstance(b, Heading))
        print(f"{scope.stem}: {len(scope.lessons())} lessons, {headings} headings")
        if scope.stem == "handbook-full":
            check_coverage(page, docx_path)
    for stale in DOWNLOADS.glob("handbook-*"):
        if stale.name not in keep:
            stale.unlink()
    manifest = sorted(keep)
    (DOWNLOADS / "handbook-manifest.json").write_text(json.dumps(manifest, indent=1) + "\n")


if __name__ == "__main__":
    main()
